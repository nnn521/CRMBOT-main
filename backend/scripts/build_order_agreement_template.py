"""
生成专业订车协议模板 docx

依据《中华人民共和国民法典》合同编，包含：
- 车辆预订信息（详细表格）
- 定金条款（明确定金法律含义及罚则）
- 车辆锁定与配置保护
- 价格保护（厂家调价处理）
- 预计到车日期与逾期处理
- 客户取消与退款条件
- 转单/转让条款
- 签署栏

用法:
    cd backend && python scripts/build_order_agreement_template.py
"""

import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import re as _re
# ═══════════════════════════════════════════════════════════════
# 样式工具函数（同 build_car_purchase_template.py）
# ═══════════════════════════════════════════════════════════════
FONT_TITLE, FONT_HEADING, FONT_BODY, FONT_BOLD_LABEL = "黑体", "黑体", "宋体", "黑体"
SIZE_TITLE, SIZE_HEADING, SIZE_BODY, SIZE_SMALL = Pt(18), Pt(12), Pt(10.5), Pt(9)

def set_font(run, name=FONT_BODY, size=SIZE_BODY, bold=False, color=None):
    run.font.name = name; run.font.size = size; run.bold = bold
    if color: run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="{name}"/>'); rPr.insert(0, rFonts)
    else: rFonts.set(qn('w:eastAsia'), name)

def add_paragraph(doc, text, font=FONT_BODY, size=SIZE_BODY, bold=False, alignment=None, space_after=Pt(6), space_before=Pt(0), indent=None):
    p = doc.add_paragraph()
    if alignment is not None: p.alignment = alignment
    pf = p.paragraph_format; pf.space_after = space_after; pf.space_before = space_before
    if indent is not None: pf.first_line_indent = indent
    run = p.add_run(text); set_font(run, font, size, bold)
    return p

def add_clause_heading(doc, number, title):
    p = doc.add_paragraph(); pf = p.paragraph_format; pf.space_before = Pt(12); pf.space_after = Pt(6)
    run = p.add_run(f"第{number}条  {title}"); set_font(run, FONT_HEADING, SIZE_HEADING, bold=True)
    return p

def add_placeholder_para(doc, label, placeholder, indent=Cm(0.74)):
    p = doc.add_paragraph(); pf = p.paragraph_format; pf.space_after = Pt(2)
    if indent: pf.left_indent = indent
    r1 = p.add_run(label); set_font(r1, FONT_BOLD_LABEL, SIZE_BODY, bold=True)
    r2 = p.add_run(f"{{{{{placeholder}}}}}"); set_font(r2, FONT_BODY, SIZE_BODY)
    return p

def add_body_para(doc, text, indent=Cm(0.74)):
    return add_paragraph(doc, text, FONT_BODY, SIZE_BODY, indent=indent, space_after=Pt(4))

def set_table_borders(table):
    tbl = table._tbl; tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/><w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/><w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/></w:tblBorders>')
    tblPr.append(borders)

OUTPUT_PATH = "data/templates/contracts/order_agreement.docx"


def build_template():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    # ─── 标题 ───
    add_paragraph(doc, "汽 车 订 购 协 议 书", FONT_TITLE, SIZE_TITLE, bold=True,
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(16))

    add_placeholder_para(doc, "协议编号：", "agreement_no", indent=None)
    add_paragraph(doc, "", space_after=Pt(8))

    # ═══════════════════════════════════════
    # 甲乙双方
    # ═══════════════════════════════════════
    add_clause_heading(doc, "一、", "协议当事人")
    add_placeholder_para(doc, "甲方（经销商）：", "dealer_name")
    add_placeholder_para(doc, "地址：", "dealer_address")
    add_placeholder_para(doc, "联系电话：", "dealer_phone")
    add_paragraph(doc, "", space_after=Pt(4))
    add_placeholder_para(doc, "乙方（客户）：", "buyer_name")
    add_placeholder_para(doc, "身份证号：", "buyer_id_card")
    add_placeholder_para(doc, "联系电话：", "buyer_phone")
    add_placeholder_para(doc, "地址：", "buyer_address")

    # ═══════════════════════════════════════
    # 第二条 预订车辆信息
    # ═══════════════════════════════════════
    add_clause_heading(doc, "二", "预订车辆信息")
    add_body_para(doc, '乙方自愿向甲方预订以下车辆（以下简称"预订车辆"），甲方同意接受预订：')

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    hdr_cells = table.rows[0].cells
    for i, text in enumerate(["项  目", "内  容"]):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(text)
        set_font(run, FONT_BOLD_LABEL, SIZE_SMALL, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    order_fields = [
        ("汽车品牌", "car_brand"),
        ("型号规格", "car_model"),
        ("配置/版本", "car_config"),
        ("车身颜色（外观）", "car_color"),
        ("内饰颜色", "interior_color"),
        ("排量/发动机类型", "displacement"),
        ("生产厂名称", "manufacturer"),
        ("市场指导价（元）", "estimated_price"),
    ]
    for label, placeholder in order_fields:
        row = table.add_row()
        c0 = row.cells[0]; c0.text = ""
        p0 = c0.paragraphs[0]; r0 = p0.add_run(label)
        set_font(r0, FONT_BOLD_LABEL, SIZE_SMALL, bold=True)
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        c1 = row.cells[1]; c1.text = ""
        p1 = c1.paragraphs[0]; r1 = p1.add_run(f"{{{{{placeholder}}}}}")
        set_font(r1, FONT_BODY, SIZE_SMALL)
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)

    add_paragraph(doc, "", space_after=Pt(8))

    # ═══════════════════════════════════════
    # 第三条 定金条款
    # ═══════════════════════════════════════
    add_clause_heading(doc, "三", "定金条款")
    add_body_para(doc, "1. 乙方应于本协议签订之日起 {{deposit_days}} 日内，向甲方支付预订定金人民币 {{deposit_amount}} 元（大写：{{deposit_amount_cn}} 元整）。")
    add_body_para(doc, '2. 本协议中的"定金"为《中华人民共和国民法典》第五百八十七条规定的违约定金，具有担保合同履行的法律效力，而非预付款性质的"订金"。')
    add_body_para(doc, "3. 定金金额不超过预订车辆市场指导价的 20%。")
    add_body_para(doc, "4. 定金支付方式：{{deposit_payment_method}}。甲方收款账户：{{dealer_account}}。")

    # ═══════════════════════════════════════
    # 第四条 车辆锁定与配置保护
    # ═══════════════════════════════════════
    add_clause_heading(doc, "四", "车辆锁定与配置保护")
    add_body_para(doc, "1. 甲方收到定金后，应立即为乙方锁定上述车型、配置及颜色，锁定期间为 {{lock_days}} 日。在锁定期内，甲方不得将该车辆另行出售或预订给第三方。")
    add_body_para(doc, "2. 在锁定期内，乙方如要求变更颜色或配置，应在车辆排产前书面通知甲方。如已排产，变更所产生的额外费用由乙方承担。")
    add_body_para(doc, "3. 甲方应确保预订车辆为全新商品车，非展车、非试驾车、非事故车、非泡水车，出厂日期不早于交付日期 6 个月。")

    # ═══════════════════════════════════════
    # 第五条 价格保护
    # ═══════════════════════════════════════
    add_clause_heading(doc, "五", "价格保护条款")
    add_body_para(doc, "1. 预订车辆最终成交价以提车时签订正式《汽车买卖合同》为准。双方约定的优惠方式为：{{price_protection_method}}。")
    add_body_para(doc, '2. 如提车时厂家指导价下调，甲方应按下调后的价格执行。如遇厂家调涨指导价，按本协议签订时的厂家指导价执行（即"就低不就高"原则）。')
    add_body_para(doc, "3. 双方可就额外精品、保养套餐、保险等另行协商，并写入正式合同。")

    # ═══════════════════════════════════════
    # 第六条 到车与交付
    # ═══════════════════════════════════════
    add_clause_heading(doc, "六", "到车与交付")
    add_placeholder_para(doc, "1. 预计到车日期：", "eta_date")
    add_body_para(doc, "2. 甲方应在预订车辆到店后 3 个工作日内通知乙方前来验车，通知方式包括电话、短信或微信。")
    add_body_para(doc, "3. 乙方应在收到通知后 {{acceptance_days}} 日内到店验车并签署正式《汽车买卖合同》。逾期未到的，甲方有权从第 {{acceptance_days}} 日后按每日 {{late_acceptance_fee}} 元收取车辆保管费。")
    add_body_para(doc, "4. 验车时，双方按照本协议第二条约定逐项核对车辆信息。如有不符，乙方有权拒绝接收，甲方应重新提供符合约定的车辆。")

    # ═══════════════════════════════════════
    # 第七条 逾期处理
    # ═══════════════════════════════════════
    add_clause_heading(doc, "七", "逾期未到车处理")
    add_body_para(doc, "1. 如预计到车日期届满后 {{grace_period_days}} 日内车辆仍未到店，乙方有权选择：")
    add_body_para(doc, "   (1) 继续等待，双方重新约定交车时间并签署补充协议；")
    add_body_para(doc, "   (2) 解除本协议，要求甲方双倍返还已收定金。")
    add_body_para(doc, "2. 因厂家停产、自然灾害、政府政策调整等不可抗力导致无法交付的，甲方不承担违约责任，但应将已收定金无息全额退还乙方。甲方应提供厂家停产通知等相关证明材料。")
    add_body_para(doc, "3. 因物流运输延误、厂家产能不足等非不可抗力原因导致逾期的，甲方应按每逾期一日向乙方支付定金金额 {{overdue_rate}}% 的违约金。")

    # ═══════════════════════════════════════
    # 第八条 客户取消与退款
    # ═══════════════════════════════════════
    add_clause_heading(doc, "八", "客户取消预订与退款")
    add_body_para(doc, "1. 在车辆排产前，乙方如主动取消预订，甲方扣除定金 {{cancel_deduction_rate}}%（不超过50%）后将余款退还乙方。")
    add_body_para(doc, "2. 车辆排产后、到车前，乙方如主动取消预订，甲方不予退还定金，用于弥补库存调配损失。")
    add_body_para(doc, "3. 如因甲方原因（包括但不限于第七条所列情形）导致无法按约定交付车辆的，乙方有权取消预订并要求甲方双倍返还定金。")
    add_body_para(doc, "4. 如因贷款审批未通过导致乙方无法继续购车的，双方均有权解除本协议，甲方应将定金全额退还乙方，互不承担违约责任。")

    # ═══════════════════════════════════════
    # 第九条 转单条款
    # ═══════════════════════════════════════
    add_clause_heading(doc, "九", "预订名额转让")
    add_body_para(doc, "1. 乙方在签署正式《汽车买卖合同》前，经甲方书面同意，可将本协议项下的预订名额转让给第三方。")
    add_body_para(doc, "2. 转让时，第三方应与甲方重新签署预订协议，原协议自动终止。乙方向第三方转让所产生的费用由乙方自行承担。")
    add_body_para(doc, "3. 转让手续费：{{transfer_fee}} 元（如需），由乙方承担。")

    # ═══════════════════════════════════════
    # 第十条 争议解决与其他
    # ═══════════════════════════════════════
    add_clause_heading(doc, "十", "争议解决与其他约定")
    add_body_para(doc, "1. 本协议在履行中发生争议的，双方应先友好协商；协商不成的，可向甲方所在地人民法院提起诉讼。")
    add_body_para(doc, "2. 甲方对知悉的乙方个人信息负有保密义务。")
    add_body_para(doc, "3. 双方签署正式《汽车买卖合同》后，本协议即行终止，双方权利义务以正式合同为准。")
    add_body_para(doc, "4. 本协议一式贰份，双方各执壹份，自双方签字或盖章之日起生效。")
    add_body_para(doc, "5. 本协议未尽事宜，双方可另行签订补充协议，补充协议与本协议具有同等法律效力。")

    add_paragraph(doc, "", space_after=Pt(20))

    # ═══════════════════════════════════════
    # 签署栏
    # ═══════════════════════════════════════
    sign_table = doc.add_table(rows=1, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(sign_table)
    for i, (party, name_key) in enumerate([
        ("甲 方（经销商盖章）", "dealer_name"),
        ("乙 方（客户签字）", "buyer_name"),
    ]):
        cell = sign_table.rows[0].cells[i]
        cell.text = ""
        p1 = cell.paragraphs[0]
        r1 = p1.add_run(f"{party}\n\n\n")
        set_font(r1, FONT_BOLD_LABEL, SIZE_BODY, bold=True)
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2 = cell.add_paragraph()
        r2 = p2.add_run("签署日期：           年     月     日")
        set_font(r2, FONT_BODY, SIZE_SMALL)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(f"OK: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_template()
