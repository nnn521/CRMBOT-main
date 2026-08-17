"""
生成正式企业车辆采购方案书模板 docx

参照企业商务提案标准格式：
- 执行摘要
- 项目背景与需求分析
- 推荐车型方案（对比表）
- 竞品对标分析
- 金融方案与费用明细
- 售后服务与增值项目
- 实施计划与风险提示
- 商务条款与联系方式

去 AI 味：正式商务用语，客观事实陈述，完整表格结构
用法: cd backend && python scripts/build_proposal_template.py
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ═══════════════════════════════════════════
FONT_TITLE, FONT_HEADING, FONT_BODY, FONT_BOLD = "黑体", "黑体", "宋体", "黑体"
SIZE_TITLE, SIZE_H1, SIZE_H2, SIZE_BODY, SIZE_SMALL = Pt(22), Pt(14), Pt(12), Pt(10.5), Pt(9)
BLUE = RGBColor(0x1A,0x3C,0x5E)
GRAY = RGBColor(0x66,0x66,0x66)

def set_font(run, name=FONT_BODY, size=SIZE_BODY, bold=False, color=None):
    run.font.name = name; run.font.size = size; run.bold = bold
    if color: run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="{name}"/>'); rPr.insert(0, rFonts)
    else: rFonts.set(qn('w:eastAsia'), name)

def add_p(doc, text, font=FONT_BODY, size=SIZE_BODY, bold=False, align=None, sa=Pt(6), sb=Pt(0), indent=None, color=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    pf = p.paragraph_format; pf.space_after = sa; pf.space_before = sb
    if indent: pf.left_indent = indent
    r = p.add_run(text); set_font(r, font, size, bold, color)
    return p

def add_sec(doc, title):
    """章节标题：带分隔线效果"""
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(20); pf.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="4" w:color="1A3C5E"/></w:pBdr>')
    pPr.append(pBdr)
    r = p.add_run(title); set_font(r, FONT_HEADING, SIZE_H2, bold=True, color=BLUE)
    return p

def add_ph(doc, label, placeholder, indent=Cm(0.74)):
    p = doc.add_paragraph(); pf = p.paragraph_format; pf.space_after = Pt(2)
    if indent: pf.left_indent = indent
    r1 = p.add_run(label); set_font(r1, FONT_BOLD, SIZE_BODY, bold=True)
    r2 = p.add_run(f"{{{{{placeholder}}}}}"); set_font(r2, FONT_BODY, SIZE_BODY)
    return p

def add_body(doc, text, indent=Cm(0.74)):
    return add_p(doc, text, FONT_BODY, SIZE_BODY, indent=indent, sa=Pt(4))

def set_tbl_borders(table):
    tbl = table._tbl; tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="333333"/><w:left w:val="single" w:sz="4" w:space="0" w:color="333333"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="333333"/><w:right w:val="single" w:sz="4" w:space="0" w:color="333333"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="333333"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="333333"/></w:tblBorders>')
    tblPr.append(borders)

def make_table(doc, headers, rows_data):
    """创建格式化表格：headers 列表 + rows_data（每行是键值对列表）"""
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; set_tbl_borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); set_font(r, FONT_BOLD, SIZE_SMALL, bold=True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        # 表头底色
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1A3C5E" w:val="clear"/>')
        c._tc.get_or_add_tcPr().append(shading)
        r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    for row_data in rows_data:
        row = t.add_row()
        for i, text in enumerate(row_data):
            c = row.cells[i]; c.text = ""
            rn = c.paragraphs[0].add_run(str(text))
            set_font(rn, FONT_BOLD if i==0 else FONT_BODY, SIZE_SMALL, bold=(i==0))
            c.paragraphs[0].paragraph_format.space_before = Pt(1)
            c.paragraphs[0].paragraph_format.space_after = Pt(1)
            if i > 0: c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return t

OUTPUT_PATH = "data/templates/proposals/proposal.docx"


def build_template():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.54); section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8); section.right_margin = Cm(2.8)

    # ═══════════════════════════════════════════════
    # 封面
    # ═══════════════════════════════════════════════
    add_p(doc, "", sb=Pt(48))
    add_p(doc, "企 业 车 辆 采 购 方 案 书", FONT_TITLE, SIZE_TITLE, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, sa=Pt(12), color=BLUE)
    add_p(doc, "CORPORATE VEHICLE PROCUREMENT PROPOSAL", FONT_BODY, Pt(10),
          align=WD_ALIGN_PARAGRAPH.CENTER, sa=Pt(32), color=GRAY)
    add_p(doc, "— 商 务 方 案 · 内 部 资 料 —", FONT_BODY, Pt(10),
          align=WD_ALIGN_PARAGRAPH.CENTER, sa=Pt(24), color=GRAY)

    # 文稿信息
    make_table(doc, ["项  目", "内  容"], [
        ["方案编号", "{{proposal_no}}"],
        ["呈报单位", "{{seller_name}}"],
        ["致", "{{customer_company}}"],
        ["联系人", "{{contact_person}}"],
        ["联系电话", "{{contact_phone}}"],
        ["项目名称", "{{project_name}}"],
        ["编制日期", "{{proposal_date}}"],
        ["方案有效期", "{{proposal_valid_days}}个工作日"],
    ])

    add_p(doc, "", sa=Pt(12))
    add_p(doc, "本方案书仅供呈报单位与受文单位之间内部使用，未经双方书面同意，不得向第三方披露。",
          FONT_BODY, Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY)

    doc.add_page_break()

    # ═══════════════════════════════════════════════
    # 一、执行摘要
    # ═══════════════════════════════════════════════
    add_sec(doc, "一、执行摘要")
    add_body(doc, "本文档系呈报单位为 {{customer_company}} 编制的企业车辆采购方案书，旨在为贵公司公务/商务用车需求提供系统化解决方案。")
    add_body(doc, "经初步需求调研及市场分析，本方案推荐以 {{car_model_1}} 为主选车型，{{car_model_2}} 为备选车型，并提供相应的金融方案、费用明细及售后服务计划。")
    add_body(doc, "方案要点如下：")
    add_body(doc, "（一）主推车型：{{car_model_1}}，厂商指导价 {{price_1}}，综合油耗 {{fuel_1}} L/100km。")
    add_body(doc, "（二）采购预算：单车落地价约 {{landing_total}} 元（含购置税、保险、上牌费）。")
    add_body(doc, "（三）交付周期：预计合同签订后 {{estimated_delivery_weeks}} 周内完成交付。")
    add_body(doc, "（四）售后服务：包含 {{maintenance_plan}}，{{extended_warranty}}。")

    # ═══════════════════════════════════════════════
    # 二、项目背景与需求分析
    # ═══════════════════════════════════════════════
    add_sec(doc, "二、项目背景与需求分析")
    add_body(doc, "（一）项目背景")
    add_body(doc, "{{company_intro}}")
    add_body(doc, "（二）用车需求")
    make_table(doc, ["需求维度", "具体内容"], [
        ["用车场景", "{{usage_scenario}}"],
        ["购车预算", "{{budget_range}}"],
        ["预计采购数量", "{{fleet_size}}"],
        ["车型偏好", "{{car_preference}}"],
        ["重点关注", "{{priority_concerns}}"],
    ])
    add_body(doc, "（三）需求分析结论")
    add_body(doc, "{{needs_analysis}}")

    # ═══════════════════════════════════════════════
    # 三、推荐车型方案
    # ═══════════════════════════════════════════════
    add_sec(doc, "三、推荐车型方案")
    add_body(doc, "根据上述需求，经对当前市场在售车型进行筛选与对比，推荐以下方案：")
    add_body(doc, "方案一：{{car_model_1}}（主推方案）", indent=None)

    make_table(doc, ["技术参数", "规格数据"], [
        ["品牌及型号", "{{car_model_1}}"],
        ["厂商指导价", "{{price_1}}"],
        ["发动机/排量", "{{displacement_1}}"],
        ["变速箱类型", "{{transmission_1}}"],
        ["综合油耗（L/100km）", "{{fuel_1}}"],
        ["车身尺寸（长×宽×高mm）", "{{size_1}}"],
        ["轴距（mm）", "{{wheelbase_1}}"],
        ["推荐理由", "{{reason_1}}"],
    ])

    add_p(doc, "", sa=Pt(10))
    add_body(doc, "方案二：{{car_model_2}}（备选方案）", indent=None)

    make_table(doc, ["技术参数", "规格数据"], [
        ["品牌及型号", "{{car_model_2}}"],
        ["厂商指导价", "{{price_2}}"],
        ["发动机/排量", "{{displacement_2}}"],
        ["变速箱类型", "{{transmission_2}}"],
        ["综合油耗（L/100km）", "{{fuel_2}}"],
        ["车身尺寸（长×宽×高mm）", "{{size_2}}"],
        ["轴距（mm）", "{{wheelbase_2}}"],
        ["推荐理由", "{{reason_2}}"],
    ])

    # ═══════════════════════════════════════════════
    # 四、竞品对标分析
    # ═══════════════════════════════════════════════
    add_sec(doc, "四、竞品对标分析")
    add_body(doc, "以下为推荐方案与主要竞品车型的关键指标对比：")
    add_body(doc, "{{comparison_summary}}")

    make_table(doc, ["指标", "{{car_model_1}}", "{{car_model_2}}", "{{competitor_1}}", "{{competitor_2}}", "{{competitor_3}}"], [
        ["厂商指导价", "{{cmp_price_1}}", "{{cmp_price_2}}", "{{cmp_price_3}}", "{{cmp_price_4}}", "{{cmp_price_5}}"],
        ["综合油耗", "{{cmp_fuel_1}}", "{{cmp_fuel_2}}", "{{cmp_fuel_3}}", "{{cmp_fuel_4}}", "{{cmp_fuel_5}}"],
        ["3年保值率", "{{cmp_resale_1}}", "{{cmp_resale_2}}", "{{cmp_resale_3}}", "{{cmp_resale_4}}", "{{cmp_resale_5}}"],
        ["轴距（mm）", "{{cmp_wheelbase_1}}", "{{cmp_wheelbase_2}}", "{{cmp_wheelbase_3}}", "{{cmp_wheelbase_4}}", "{{cmp_wheelbase_5}}"],
        ["核心优势", "{{cmp_advantage_1}}", "{{cmp_advantage_2}}", "{{cmp_advantage_3}}", "{{cmp_advantage_4}}", "{{cmp_advantage_5}}"],
    ])

    # ═══════════════════════════════════════════════
    # 五、金融方案
    # ═══════════════════════════════════════════════
    add_sec(doc, "五、金融方案")
    add_body(doc, "针对本次采购，我司提供以下三种融资方案供贵公司选择：")
    add_body(doc, "{{finance_plan}}")

    make_table(doc, ["方案", "首付比例", "首付金额", "月供金额", "贷款期限"], [
        ["方案一", "{{plan1_ratio}}", "{{plan1_down}}", "{{plan1_monthly}}", "{{plan1_term}}"],
        ["方案二", "{{plan2_ratio}}", "{{plan2_down}}", "{{plan2_monthly}}", "{{plan2_term}}"],
        ["方案三", "{{plan3_ratio}}", "{{plan3_down}}", "{{plan3_monthly}}", "{{plan3_term}}"],
    ])
    add_body(doc, "注：以上月供金额为预估计算值，最终以金融机构审批结果为准。贷款利率按放款日全国银行间同业拆借中心公布的同期LPR为基准，上下浮动以银行批复为准。")

    # ═══════════════════════════════════════════════
    # 六、落地费用明细
    # ═══════════════════════════════════════════════
    add_sec(doc, "六、落地费用明细")
    add_body(doc, "以推荐方案一（{{car_model_1}}）为测算基础，单车预估落地费用如下：")

    make_table(doc, ["费用项目", "金额（元）", "计算依据"], [
        ["裸车价", "{{total_car_price}}", "厂商指导价，最终以合同为准"],
        ["车辆购置税", "{{purchase_tax}}", "裸车价÷1.13×10%，以税务局核定为准"],
        ["交强险+商业险", "{{insurance_fee}}", "首年预估，含车损险/三者险/不计免赔"],
        ["上牌费", "{{plate_fee}}", "含验车、牌照、行驶证工本费"],
        ["车船税", "{{vessel_tax}}", "按排量及当地标准征收"],
        ["预估落地总价", "{{landing_total}}", "以上各项合计"],
    ])
    add_body(doc, "{{cost_breakdown}}")

    # ═══════════════════════════════════════════════
    # 七、售后服务与增值项目
    # ═══════════════════════════════════════════════
    add_sec(doc, "七、售后服务与增值项目")
    add_body(doc, "（一）随车精品礼包")
    add_body(doc, "{{gift_package}}（参考价值约 {{gift_value}}）。")
    add_body(doc, "（二）保养服务")
    add_body(doc, "{{maintenance_plan}}。")
    add_body(doc, "（三）延长保修")
    add_body(doc, "{{extended_warranty}}。")
    add_body(doc, "（四）其他增值服务")
    add_body(doc, "{{value_added_services}}")

    # ═══════════════════════════════════════════════
    # 八、实施计划与交付
    # ═══════════════════════════════════════════════
    add_sec(doc, "八、实施计划与交付")
    add_body(doc, "（一）实施时间表")
    make_table(doc, ["阶段", "工作内容", "预计耗时"], [
        ["需求确认", "双方确认车型、配置、颜色、数量", "3-5个工作日"],
        ["商务洽商", "价格谈判、金融方案确定、合同签署", "5-7个工作日"],
        ["车辆调配", "厂家排产/调拨、物流运输", "{{estimated_delivery_weeks}}"],
        ["交付验收", "到店验收、上牌、交车仪式", "3-5个工作日"],
    ])
    add_body(doc, "（二）交付标准")
    add_body(doc, "1. 车辆须为出厂日期不早于交付日6个月的全新商品车，非展车、非试驾车、非事故车。")
    add_body(doc, "2. 交付时提供：购车发票、车辆合格证、三包凭证、使用说明书、保养手册、PDI检测报告。")
    add_body(doc, "3. 交车时由双方共同验收，签署《商品车辆交接书》。")

    # ═══════════════════════════════════════════════
    # 九、风险提示与免责声明
    # ═══════════════════════════════════════════════
    add_sec(doc, "九、风险提示与免责声明")
    add_body(doc, "1. 以上价格为当前厂家指导价，实际成交价以签订正式购车合同为准。如遇厂家调整指导价，按合同约定执行。")
    add_body(doc, "2. 金融方案中的利率及月供金额为参考测算值，最终以金融机构审批结果为准。")
    add_body(doc, "3. 落地费用中的购置税、保险费为预估金额，实际金额以当地税务及保险机构核定为准。")
    add_body(doc, "4. 车辆交付时间受厂家产能、物流运输等因素影响，以上时间为正常情况下的预估周期。")
    add_body(doc, "5. 部分城市有限行、限牌政策，请贵公司提前了解并确认购车资质。")
    add_body(doc, "6. 年度用车成本估算：以年行驶 {{annual_mileage}} 公里计，燃油费约 {{annual_fuel_cost}} 元，保养费约 {{annual_maintenance_cost}} 元。")
    add_body(doc, "{{risk_disclaimer}}")

    # ═══════════════════════════════════════════════
    # 十、商务联系
    # ═══════════════════════════════════════════════
    add_sec(doc, "十、商务联系")
    add_body(doc, "如贵公司对本方案有任何疑问或需进一步洽谈，请通过以下方式联系：")

    make_table(doc, ["项目", "信息"], [
        ["销售顾问", "{{sales_rep}}"],
        ["联系电话", "{{sales_phone}}"],
        ["电子邮箱", "{{sales_email}}"],
        ["展厅地址", "{{showroom_address}}"],
    ])

    # ═══════════════════════════════════════════════
    # 尾页
    # ═══════════════════════════════════════════════
    add_p(doc, "", sa=Pt(32))
    add_p(doc, "— 本方案书依据实际市场数据编制，仅供参考。最终商务条款以双方签订的书面合同为准。—",
          FONT_BODY, Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY)
    add_p(doc, "© {{seller_name}}  {{proposal_date}}",
          FONT_BODY, Pt(8), align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY)

    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(f"OK: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_template()
